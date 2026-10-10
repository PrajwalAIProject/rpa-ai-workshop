"""Project 3 - Basic: check the Excel file your PhoneDeals bot wrote.

Reference solution. In the workshop you ask Kiro to write this file (see README.md,
prompt 4) and to run it after every bot run. It is the CHECK step of the agent loop:
never trust a bot's output without checking it.

Checks:
  1. The file and the sheet exist, with Name and Price columns.
  2. There is at least one phone.
  3. Every price is a whole number between 1 and the limit (default 20,000).
  4. No phone name is empty, and no phone is listed twice.
  5. The list is sorted by price, cheapest first.

Run:
    python check_phones.py                                  # PhonesUnder20K.xlsx, sheet Phones
    python check_phones.py C:\\PhoneDeals\\PhonesUnder20K.xlsx
    python check_phones.py PhonesUnder20K.xlsx --max 15000
"""

import sys
from pathlib import Path

from openpyxl import load_workbook

SHEET = "Phones"


def to_price(value):
    """Turn 9499, 9499.0, "9,499" or "₹9,499" into 9499; return None if it is not a number."""
    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else None
    text = str(value or "").replace("₹", "").replace(",", "").strip()
    return int(text) if text.isdigit() else None


def check(path: Path, max_price: int) -> list[str]:
    """Return a list of problems (empty list = all checks passed)."""
    if not path.exists():
        return [f"File not found: {path}. Did the bot run, and is the path right?"]
    book = load_workbook(path, read_only=True, data_only=True)
    if SHEET not in book.sheetnames:
        return [f"Sheet '{SHEET}' not found. Sheets in the file: {', '.join(book.sheetnames)}"]
    rows = list(book[SHEET].iter_rows(values_only=True))
    if not rows:
        return ["The sheet is empty."]
    header = [str(h or "").strip().lower() for h in rows[0]]
    if "name" not in header or "price" not in header:
        return [f"Need columns 'Name' and 'Price' in row 1; found: {rows[0]}"]
    name_col, price_col = header.index("name"), header.index("price")

    problems, seen, prices = [], set(), []
    data = [r for r in rows[1:] if any(c not in (None, "") for c in r)]
    if not data:
        problems.append("No phones in the sheet: the bot found nothing, or did not write the rows.")
    for line, row in enumerate(data, start=2):
        name = str(row[name_col] or "").strip()
        price = to_price(row[price_col])
        if not name:
            problems.append(f"Row {line}: the phone name is empty.")
        elif name.lower() in seen:
            problems.append(f"Row {line}: '{name[:50]}' is listed twice.")
        seen.add(name.lower())
        if price is None:
            problems.append(f"Row {line}: price '{row[price_col]}' is not a whole number.")
        elif not 1 <= price <= max_price:
            problems.append(f"Row {line}: price {price:,} is above the {max_price:,} limit (or zero).")
        else:
            prices.append(price)
    if prices and prices != sorted(prices):
        problems.append("The phones are not sorted by price, cheapest first.")
    if not problems:
        print(f"  {len(data)} phones, cheapest ₹{min(prices):,}, dearest ₹{max(prices):,}")
    return problems


def main() -> int:
    args = sys.argv[1:]
    max_price = 20000
    if "--max" in args:
        i = args.index("--max")
        max_price = int(args[i + 1])
        del args[i:i + 2]
    path = Path(args[0]) if args else Path("PhonesUnder20K.xlsx")

    print(f"Checking {path} (limit ₹{max_price:,})")
    problems = check(path, max_price)
    if problems:
        print(f"FAILED: {len(problems)} problem(s)")
        for p in problems[:20]:
            print("  - " + p)
        return 1
    print("PASSED: all checks OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
